import io
import streamlit as st
import google.generativeai as genai
from pptx import Presentation
from pptx.util import Inches, Pt

# --- CẤU HÌNH TRANG WEB ---
st.set_page_config(
    page_title="Trợ Lý AI Chuyển Đổi Giáo Án Sang PowerPoint",
    page_icon="📊",
    layout="wide",
)

st.title("📊 Trợ Lý AI Chuyển Đổi Kế Hoạch Bài Dạy Sang PowerPoint (.pptx)")
st.caption("Bài tập 3: Tự động hóa thiết kế slide giảng dạy chuẩn sư phạm")

# --- LẤY API KEY ẨN TỪ SECRETS CỦA STREAMLIT CLOUD ---
try:
    GOOGLE_API_KEY = st.secrets["GOOGLE_API_KEY"]
except Exception:
    GOOGLE_API_KEY = None

# --- GIAO DIỆN NHẬP LIỆU ---
with st.form("ppt_form"):
    lesson_content = st.text_area(
        "📝 Dán nội dung Kế hoạch bài dạy (Giáo án) vào đây:",
        placeholder="Dán nội dung mục tiêu, hoạt động, nội dung bài học vào đây...",
        height=250
    )
    
    col1, col2 = st.columns(2)
    with col1:
        slide_style = st.selectbox(
            "🎨 Phong cách Slide",
            ["Hiện đại, Tối giản", "Trang trọng, Sư phạm", "Sáng tạo, Trẻ trung"]
        )
    with col2:
        num_slides = st.slider("📑 Số lượng slide ước tính:", min_value=3, max_value=10, value=5)

    submit_btn = st.form_submit_button("🚀 Tạo Tệp PowerPoint (.pptx)", type="primary")

# --- HÀM TẠO FILE POWERPOINT THỰC THỤ (.PPTX) ---
def create_powerpoint(ai_slides_text):
    prs = Presentation()
    slides_data = ai_slides_text.split("---SLIDE---")
    
    for s_data in slides_data:
        if not s_data.strip():
            continue
        
        slide_layout = prs.slide_layouts[1] # Title and Content layout
        slide = prs.slides.add_slide(slide_layout)
        
        lines = [line.strip() for line in s_data.split("\n") if line.strip()]
        title_text = "Bài Giảng Điện Tử"
        bullet_points = []
        
        for line in lines:
            if line.startswith("#") or line.startswith("Tiêu đề:") or line.startswith("Slide:"):
                title_text = line.replace("#", "").replace("Tiêu đề:", "").replace("Slide:", "").strip()
            else:
                cleaned_bullet = line.replace("-", "").replace("*", "").strip()
                if cleaned_bullet:
                    bullet_points.append(cleaned_bullet)
        
        title_shape = slide.shapes.title
        title_shape.text = title_text
        
        if slide.placeholders and len(slide.placeholders) > 1:
            body_shape = slide.placeholders[1]
            tf = body_shape.text_frame
            tf.text = ""
            
            for i, bp in enumerate(bullet_points):
                if i == 0:
                    p = tf.paragraphs[0]
                else:
                    p = tf.add_paragraph()
                p.text = bp
                p.level = 0
                p.font.size = Pt(20)

    pptx_buffer = io.BytesIO()
    prs.save(pptx_buffer)
    pptx_buffer.seek(0)
    return pptx_buffer

# --- XỬ LÝ KHI BẤM NÚT ---
if submit_btn:
    if not GOOGLE_API_KEY:
        st.error("⚠️ Hệ thống chưa cấu hình API Key ở mục Secrets trên Streamlit Cloud!")
    elif not lesson_content:
        st.error("⚠️ Vui lòng dán nội dung kế hoạch bài dạy vào ô trống!")
    else:
        try:
            genai.configure(api_key=GOOGLE_API_KEY)
            model = genai.GenerativeModel("gemini-2.5-flash")

            prompt = f"""
            Bạn là trợ lý chuyên gia thiết kế bài giảng điện tử. 
            Dựa trên nội dung Kế hoạch bài dạy sau đây, hãy cô đọng và cấu trúc thành đúng {num_slides} Slide PowerPoint giảng dạy rõ ràng, xúc tích, bám sát giáo án.

            NỘI DUNG GIÁO ÁN:
            {lesson_content}

            YÊU CẦU ĐỊNH DẠNG:
            Mỗi slide phải bắt đầu bằng từ khóa `---SLIDE---`, theo sau là tiêu đề và các ý chính (dùng dấu gạch đầu dòng `-`). Không viết văn bản dài dòng, chỉ lấy các ý trọng tâm.
            """

            with st.spinner("🔄 AI đang phân tích và thiết kế cấu trúc slide..."):
                response = model.generate_content(prompt)
                ai_output = response.text

            st.success("✅ Đã thiết kế slide thành công!")
            st.markdown("---")
            st.markdown(ai_output)

            ppt_file = create_powerpoint(ai_output)
            
            st.markdown("---")
            st.download_button(
                label="📥 Tải xuống Tệp PowerPoint (.pptx) để nộp bài",
                data=ppt_file,
                file_name="BaiGiang_PowerPoint.pptx",
                mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                type="primary"
            )

        except Exception as e:
            st.error(f"❌ Có lỗi xảy ra: {str(e)}")